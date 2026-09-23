"""
SmartAttend AI - Word Documentation (.docx) Generator
Generates a comprehensive, professional project documentation report.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tc_pr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal margins (padding) for a table cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tc_pr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Applies subtle borders to a table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, text, title="NOTE:", border_color="2563EB", bg_color="F0F4FF"):
    """Creates a stylized callout box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tc_pr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"📌 {title} ")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    doc.add_paragraph() # Spacing

def add_code_block(doc, code_str):
    """Adds a formatted code snippet box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tc_pr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    doc.add_paragraph() # Spacing

def format_heading(p, text, level=1):
    """Formats headings with beautiful typography."""
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x0F, 0x29, 0x5A) # Deep Royal Navy
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Navy Blue
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Vivid Blue

def add_body_p(doc, text="", bold_prefix=None, space_after=4):
    """Adds standard body paragraph with high legibility."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
        r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
    if text:
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    """Adds a bullet point with custom styling."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def build_styled_table(doc, headers, data, col_widths=None):
    """Generates a professional styled table with zebra striping and custom column widths."""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for idx, title in enumerate(headers):
        hdr_cells[idx].text = title
        set_cell_background(hdr_cells[idx], "0F295A") # Dark Navy
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                
    # Set widths if specified
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    set_table_borders(table, color="CBD5E1", sz="4")
    doc.add_paragraph() # Spacing
    return table

print("Helper definitions loaded successfully.")
